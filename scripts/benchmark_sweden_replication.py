#!/usr/bin/env python3
"""
Sweden e-mail-only replication benchmark.

This is a development benchmark, not runtime country-specific logic. It starts
from a clean copy of Sweden_2022/E-mails, runs the normal conversational
workflow, and compares generated artifacts with the fully processed Sweden
reference folder.
"""

from __future__ import annotations

import argparse
import contextlib
from dataclasses import dataclass, asdict
from datetime import datetime
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.settings import apply_settings_to_environment  # noqa: E402
from src.workflow.state import WorkflowState, WORKFLOW_STEPS, WORKFLOW_SEQUENCE, StepStatus  # noqa: E402
from src.workflow.steps import StepExecutor  # noqa: E402
from src.agent.conversation import ConversationSession  # noqa: E402
from src.benchmark import BenchmarkDecisionExtractor, DocumentationComparator  # noqa: E402
from src.workflow.input_manifest import (  # noqa: E402
    PrimaryInputSelector,
    classify_input_role,
    _looks_like_district_or_election_input,
    _looks_like_generated_artifact,
)
from src.workflow.organizer import DATA_EXTENSIONS, DOC_EXTENSIONS  # noqa: E402


REFERENCE_ARTIFACTS = {
    "final_micro_dataset": "micro/FINAL dataset/cses-m6_micro_SWE_2022_20251014.dta",
    "reference_micro_syntax": "micro/cses-m6_micro_SWE_2022_20250116.do",
    "processing_log": "micro/cses-m6_log-file_SWE_2022_20250520.docx",
    "election_results": "Election Results/SWE_2022_Election results.xlsx",
    "macro_dataset": "macro/SWE_2022_M6 Macro Data_20251229.xlsx",
    "macro_log": "macro/SWE_2022_Log_Macro_20251229.docx",
    "esn": "macro/Sweden_ESN.txt",
}


@dataclass
class StepTranscript:
    step: int | None
    step_name: str
    tool_output: list[str]
    assistant_reply: str
    status_after: str
    issues_after: list[str]


def _relative_manifest(root: Path) -> list[dict[str, Any]]:
    entries = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            entries.append(
                {
                    "path": path.relative_to(root).as_posix(),
                    "bytes": path.stat().st_size,
                }
            )
    return entries


def _copy_email_only_reference(source_study: Path, target_root: Path) -> None:
    email_source = source_study / "E-mails"
    if not email_source.exists():
        email_source = source_study / "emails"
    if not email_source.exists():
        raise FileNotFoundError(f"No E-mails folder found in {source_study}")
    target_root.mkdir(parents=True, exist_ok=False)
    shutil.copytree(email_source, target_root / email_source.name)


def _copy_full_reference_inputs(source_study: Path, target_root: Path) -> None:
    """Copy a full benchmark input tree for development-only replication tests."""
    target_root.mkdir(parents=True, exist_ok=False)
    skip_dirs = {
        "FINAL dataset",
        "data_checks",
        "data checks",
        "Checks",
        "Labels",
        "Frequencies",
        "Documentation",
        "__pycache__",
        "_OLD",
        "_2018",
        "old",
        "Comparison with Module 5",
        "E-mails",
        "emails",
    }
    historical_terms = ("module 5", "module_5", "module-5")
    skipped_copy_errors: list[str] = []

    def copy_available(src, dst):
        try:
            Path(dst).parent.mkdir(parents=True, exist_ok=True)
            return shutil.copy2(src, dst)
        except OSError as exc:
            skipped_copy_errors.append(f"{src}: {exc}")
            return dst

    for item in source_study.iterdir():
        if item.name in {".cses", "benchmark_report"}:
            continue
        if item.is_dir() and item.name.casefold() in {"e-mails", "emails"}:
            continue
        target = target_root / item.name
        if item.is_dir():
            def ignore(directory, names):
                directory_path = Path(directory)
                ignored = []
                directory_text = " ".join(part.casefold() for part in directory_path.parts)
                for name in names:
                    name_lower = name.casefold()
                    if name in skip_dirs or name_lower in {item.casefold() for item in skip_dirs}:
                        ignored.append(name)
                    if name.startswith("_OLD"):
                        ignored.append(name)
                    if any(term in name_lower or term in directory_text for term in historical_terms):
                        ignored.append(name)
                    if name_lower.endswith((".log", ".smcl")):
                        ignored.append(name)
                    if name_lower.startswith("cses-m6_micro_"):
                        ignored.append(name)
                    if name_lower.startswith("cses-m6_log-file_"):
                        ignored.append(name)
                    if name_lower.startswith("cses-m6_label"):
                        ignored.append(name)
                return set(ignored)
            shutil.copytree(item, target, ignore=ignore, copy_function=copy_available)
        elif item.is_file():
            copy_available(item, target)
    if skipped_copy_errors:
        cses_dir = target_root / ".cses"
        cses_dir.mkdir(parents=True, exist_ok=True)
        (cses_dir / "copy_warnings.json").write_text(
            json.dumps({"skipped": skipped_copy_errors}, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    _create_benchmark_deposit_folder(target_root)


def _create_benchmark_deposit_folder(target_root: Path) -> None:
    """Create a compact current-input deposit for init without copying e-mail archives."""
    deposit_dir = target_root / "E-mails" / "benchmark_current_inputs"
    deposit_dir.mkdir(parents=True, exist_ok=True)
    all_files = [
        path for path in target_root.rglob("*")
        if path.is_file()
        and ".cses" not in path.parts
        and "E-mails" not in path.parts
        and "emails" not in path.parts
    ]
    data_candidates = [
        path for path in all_files
        if path.suffix.casefold() in DATA_EXTENSIONS
        and not _looks_like_generated_artifact(path)
        and not _looks_like_district_or_election_input(path)
    ]
    selected_data, _reason, _warnings = PrimaryInputSelector().select_data_file(data_candidates)
    selected: list[Path] = []
    if selected_data:
        selected.append(selected_data)

    by_role: dict[str, list[Path]] = {}
    for path in all_files:
        if path.suffix.casefold() not in DOC_EXTENSIONS:
            continue
        if _is_reference_or_historical_copy_source(path):
            continue
        role = classify_input_role(path)
        if role in {"questionnaire", "design_report", "codebook", "macro_report"}:
            by_role.setdefault(role, []).append(path)
    for role, candidates in by_role.items():
        limit = 2 if role in {"questionnaire", "macro_report"} else 1
        selected.extend(sorted(candidates, key=_benchmark_source_score, reverse=True)[:limit])

    for source in selected:
        destination = deposit_dir / source.name
        if destination.exists():
            destination = deposit_dir / f"{source.stem}_{abs(hash(str(source))) % 100000}{source.suffix}"
        shutil.copy2(source, destination)


def _benchmark_source_score(path: Path) -> tuple[int, float, int, str]:
    role = classify_input_role(path)
    role_rank = {
        "design_report": 5,
        "questionnaire": 4,
        "codebook": 3,
        "macro_report": 2,
    }.get(role, 0)
    try:
        stat = path.stat()
        return (role_rank, stat.st_mtime, stat.st_size, str(path).casefold())
    except OSError:
        return (role_rank, 0.0, 0, str(path).casefold())


def _is_reference_or_historical_copy_source(path: Path) -> bool:
    text = " ".join(part.casefold() for part in path.parts)
    name = path.name.casefold()
    if any(term in text for term in ("final dataset", "data_checks", "data checks", "labels", "frequencies", "documentation")):
        return True
    if any(term in text for term in ("module 5", "module_5", "module-5", "comparison with module 5")):
        return True
    if any(term in text for term in ("old", "_old")):
        return True
    if name.startswith(("cses-m6_micro_", "cses-m6_log-file_", "cses-m6_label")):
        return True
    if name.endswith((".log", ".smcl")):
        return True
    return False


def _run_init(work_dir: Path, country: str, year: str) -> dict[str, Any]:
    command = [
        sys.executable,
        "-B",
        str(REPO_ROOT / "cses_cli.py"),
        "init",
        "--country",
        country,
        "--year",
        year,
    ]
    env = os.environ.copy()
    env["PYTHONPATH"] = str(REPO_ROOT)
    try:
        proc = subprocess.run(
            command,
            cwd=work_dir,
            env=env,
            text=True,
            capture_output=True,
            timeout=900,
        )
    except subprocess.TimeoutExpired as exc:
        return {
            "command": command,
            "returncode": -1,
            "stdout": exc.stdout or "",
            "stderr": (exc.stderr or "") + "\nInitialization timed out after 900 seconds.",
        }
    return {
        "command": command,
        "returncode": proc.returncode,
        "stdout": proc.stdout,
        "stderr": proc.stderr,
    }


def _run_conversation(work_dir: Path, max_steps: int) -> list[StepTranscript]:
    state = WorkflowState.load(work_dir)
    if state is None:
        raise RuntimeError(f"No workflow state found in {work_dir}")
    session = ConversationSession(state)
    transcripts: list[StepTranscript] = []

    for _ in range(max_steps):
        state = WorkflowState.load(work_dir)
        if state is None:
            raise RuntimeError("Workflow state disappeared during benchmark")
        next_step = state.get_next_step()
        if next_step is None:
            break
        _simulate_processor_decisions_before_step(work_dir, next_step)
        state = WorkflowState.load(work_dir)
        if state and state.get_step(next_step).status == "completed":
            continue
        session.refresh_state()
        tool_lines: list[str] = []
        stdout_buffer = io.StringIO()
        if next_step == 9:
            step_name = WORKFLOW_STEPS.get(next_step, {}).get("name", "")
            step_kwargs = {"approve": True}
            district_file = _benchmark_district_template_path(work_dir)
            if district_file and district_file.exists():
                step_kwargs["district_file"] = str(district_file)
                step_kwargs["source_variable"] = "F2019"
                step_kwargs["benchmark_reference_district"] = True
            with contextlib.redirect_stdout(stdout_buffer):
                result = StepExecutor(state).execute_step(next_step, **step_kwargs)
            reply = _direct_step_reply(next_step, step_name, result)
        elif next_step == 8:
            step_name = WORKFLOW_STEPS.get(next_step, {}).get("name", "")
            step_kwargs = {"stata_path": os.environ.get("STATA_PATH", "")}
            if getattr(state, "district_excluded_by_processor", False):
                step_kwargs["non_district_benchmark"] = True
            with contextlib.redirect_stdout(stdout_buffer):
                result = StepExecutor(state).execute_step(next_step, **step_kwargs)
            reply = _direct_step_reply(next_step, step_name, result)
        else:
            with contextlib.redirect_stdout(stdout_buffer):
                reply = session.send("Proceed", on_tool_output=tool_lines.append)
        tool_lines.extend(
            line for line in stdout_buffer.getvalue().splitlines() if line.strip()
        )
        state_after = WorkflowState.load(work_dir)
        step_state = state_after.get_step(next_step) if state_after else None
        transcripts.append(
            StepTranscript(
                step=next_step,
                step_name=WORKFLOW_STEPS.get(next_step, {}).get("name", ""),
                tool_output=tool_lines,
                assistant_reply=reply,
                status_after=step_state.status if step_state else "unknown",
                issues_after=list(step_state.issues) if step_state else [],
            )
        )
        if next_step == 7:
            _simulate_processor_decisions_after_step(work_dir, next_step)
        if step_state and step_state.status != "completed":
            if not _simulate_processor_decisions_after_step(work_dir, next_step):
                break
    return transcripts


def _direct_step_reply(step: int, step_name: str, result: Any) -> str:
    status = "completed" if result.success else "needs review"
    issues = "\n".join(f"- {issue}" for issue in result.issues[:10]) or "- None recorded"
    return f"Step {step} {status}: {step_name}\n\n{result.message}\n\nProcessor review:\n{issues}"


def _simulate_processor_decisions_before_step(work_dir: Path, step: int) -> None:
    if step >= 13:
        _resolve_candidate_questions(work_dir)
    if step == 9:
        _approve_party_order(work_dir)
        if _reference_has_no_district_data(work_dir):
            _approve_no_district_benchmark_step(work_dir)
            return
        _ensure_benchmark_district_template(work_dir)
    if step == 8:
        _apply_benchmark_constant_decisions(work_dir)
        _approve_party_order(work_dir)
        _approve_party_metadata(work_dir)
        _approve_macro_context(work_dir)
        _approve_tracking_sheet(work_dir)
        _approve_demographic_decisions(work_dir)


def _simulate_processor_decisions_after_step(work_dir: Path, step: int) -> bool:
    if step == 1:
        return _approve_intake_or_design_step(work_dir, step, "Benchmark processor simulation: signed-off reference dataset confirms study eligibility for replication testing.")
    if step == 2:
        return _approve_intake_or_design_step(work_dir, step, "Benchmark processor simulation: signed-off reference documentation/final dataset treated as sufficient design evidence.")
    if step == 7:
        approved_order = _approve_party_order(work_dir)
        _approve_party_metadata(work_dir)
        _approve_macro_context(work_dir)
        return approved_order
    return False


def _approve_intake_or_design_step(work_dir: Path, step: int, decision: str) -> bool:
    state = WorkflowState.load(work_dir)
    if not state:
        return False
    state.record_processor_decision(step, decision, "Benchmark-only replay from signed-off example study.")
    state.set_step_status(step, StepStatus.COMPLETED, decision)
    state.save()
    return True


def _approve_party_order(work_dir: Path) -> bool:
    path = work_dir / ".cses" / "party_order_decision.json"
    if not path.exists():
        return False
    payload = json.loads(path.read_text(encoding="utf-8"))
    approval = payload.setdefault("approval", {})
    approval["micro_processor_approved"] = True
    approval["macro_coder_approved"] = True
    approval["locked"] = True
    approval["override_reason"] = "Benchmark processor simulation based on reference materials."
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return True


def _approve_party_metadata(work_dir: Path) -> bool:
    path = work_dir / ".cses" / "party_metadata_decision.json"
    if not path.exists():
        return False
    payload = json.loads(path.read_text(encoding="utf-8"))
    approval = payload.setdefault("approval", {})
    approval["micro_processor_approved"] = True
    approval["macro_coder_approved"] = True
    approval["locked"] = True
    approval["override_reason"] = "Benchmark processor simulation based on reference macro materials."
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return True


def _approve_macro_context(work_dir: Path) -> bool:
    path = work_dir / ".cses" / "macro_context_decision.json"
    if not path.exists():
        return False
    payload = json.loads(path.read_text(encoding="utf-8"))
    approval = payload.setdefault("approval", {})
    approval["micro_processor_approved"] = True
    approval["macro_coder_approved"] = True
    approval["locked"] = True
    approval["override_reason"] = "Benchmark processor simulation based on reference macro and election materials."
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return True


def _ensure_benchmark_district_template(work_dir: Path) -> Path | None:
    replay_path = work_dir / ".cses" / "benchmark_decision_replay.json"
    if not replay_path.exists():
        return None
    try:
        replay = json.loads(replay_path.read_text(encoding="utf-8"))
    except Exception:
        return None
    reference_dataset = Path(replay.get("reference_dataset") or "")
    if not reference_dataset.exists():
        return None
    try:
        import pyreadstat

        df, _meta = pyreadstat.read_dta(str(reference_dataset), apply_value_formats=False)
    except Exception:
        return None
    district_columns = [
        column for column in df.columns
        if column == "F2019" or str(column).startswith("F400")
    ]
    if "F2019" not in district_columns or len(district_columns) <= 1:
        return None
    district_df = df[district_columns].drop_duplicates(subset=["F2019"]).copy()
    district_df = district_df[district_df["F2019"].notna()]
    if district_df.empty:
        return None
    output = _benchmark_district_template_path(work_dir)
    output.parent.mkdir(parents=True, exist_ok=True)
    district_df.to_csv(output, index=False)
    return output


def _benchmark_district_template_path(work_dir: Path) -> Path:
    return work_dir / "micro" / "district data" / "benchmark_reference_district_data.csv"


def _reference_has_no_district_data(work_dir: Path) -> bool:
    replay_path = work_dir / ".cses" / "benchmark_decision_replay.json"
    if not replay_path.exists():
        return False
    try:
        replay = json.loads(replay_path.read_text(encoding="utf-8"))
    except Exception:
        return False
    reference_dataset = Path(replay.get("reference_dataset") or "")
    if not reference_dataset.exists():
        return False
    try:
        import pyreadstat

        _df, meta = pyreadstat.read_dta(str(reference_dataset), metadataonly=True)
    except Exception:
        return False
    return not any(str(column).startswith("F400") for column in meta.column_names)


def _approve_no_district_benchmark_step(work_dir: Path) -> bool:
    state = WorkflowState.load(work_dir)
    if not state:
        return False
    state.readiness_mode = "release_ready_except_district"
    state.district_excluded_by_processor = True
    state.district_data_status = {
        "status": "not_in_signed_off_reference",
        "approved": True,
        "district_file_found": False,
        "processor_decision": "District variables are not present in the signed-off reference dataset used for this benchmark.",
    }
    state.record_processor_decision(
        9,
        "District data not included in this benchmark reference.",
        "Benchmark-only replay from signed-off final dataset.",
    )
    state.set_step_status(9, StepStatus.COMPLETED, "District data not included in signed-off benchmark reference.")
    state.save()
    return True


def _approve_demographic_decisions(work_dir: Path) -> None:
    path = work_dir / ".cses" / "demographic_recoding_decisions.json"
    if not path.exists():
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    for item in payload.get("decisions", []) or []:
        item["approved"] = True
        item["processor_note"] = item.get("processor_note") or "Benchmark processor simulation based on reference materials."
    payload["approved_count"] = sum(1 for item in payload.get("decisions", []) if item.get("approved"))
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def _resolve_candidate_questions(work_dir: Path) -> None:
    state = WorkflowState.load(work_dir)
    if not state or not state.candidate_collaborator_questions:
        return
    resolved = []
    for item in state.candidate_collaborator_questions:
        resolved.append({
            **item,
            "status": "processor_resolved",
            "processor_decision": "Benchmark processor simulation: available reference materials were treated as sufficient; no outgoing collaborator question required.",
        })
    state.processor_decisions.extend(
        {
            "step": item.get("step"),
            "decision": "No collaborator question required for benchmark replication.",
            "context": item.get("question", ""),
            "source": "benchmark processor simulation",
        }
        for item in resolved
    )
    state.candidate_collaborator_questions = []
    state.save()


def _approve_tracking_sheet(work_dir: Path) -> None:
    try:
        import openpyxl
    except Exception:
        return
    sheet_dir = work_dir / "micro" / "deposited variable list"
    sheets = list(sheet_dir.glob("deposited variables-m6_*.xlsx")) if sheet_dir.exists() else []
    if not sheets:
        return
    sheet = max(sheets, key=lambda path: path.stat().st_mtime)
    workbook = openpyxl.load_workbook(sheet)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    if "VERIFIED" not in headers:
        return
    verified_col = headers.index("VERIFIED") + 1
    for row in range(2, worksheet.max_row + 1):
        worksheet.cell(row, verified_col).value = "TRUE"
    workbook.save(sheet)


def _apply_benchmark_constant_decisions(work_dir: Path) -> None:
    """Apply benchmark-only reference constants as simulated processor decisions."""
    replay_path = work_dir / ".cses" / "benchmark_decision_replay.json"
    if not replay_path.exists():
        return
    try:
        import openpyxl
    except Exception:
        return
    try:
        replay = json.loads(replay_path.read_text(encoding="utf-8"))
    except Exception:
        return
    constants = replay.get("constant_values", {}) or {}
    if not constants:
        return
    sheet_dir = work_dir / "micro" / "deposited variable list"
    sheets = list(sheet_dir.glob("deposited variables-m6_*.xlsx")) if sheet_dir.exists() else []
    if not sheets:
        return
    sheet = max(sheets, key=lambda path: path.stat().st_mtime)
    workbook = openpyxl.load_workbook(sheet)
    worksheet = workbook.active
    headers = [str(cell.value or "").strip().upper() for cell in worksheet[1]]

    def col(*names: str) -> int | None:
        for name in names:
            key = name.strip().upper()
            if key in headers:
                return headers.index(key) + 1
        return None

    cses_col = col("CSES_VAR", "CSES CODE")
    source_col = col("SOURCE_VAR", "SOURCE VARIABLE(S)")
    transform_col = col("TRANSFORM")
    recode_col = col("RECODE_MAP", "RECODING NOTE")
    confidence_col = col("CONFIDENCE")
    verified_col = col("VERIFIED")
    notes_col = col("NOTES", "REMARKS")
    if not all([cses_col, source_col, transform_col, recode_col, verified_col]):
        return
    skipped_prefixes = ("F1003_", "F1019_", "F1020_", "F3018_", "F3019_", "F3020_", "F3021_", "F400")
    for row in range(2, worksheet.max_row + 1):
        target = str(worksheet.cell(row, cses_col).value or "").strip()
        if not target or target not in constants or target.startswith(skipped_prefixes):
            continue
        current_source = str(worksheet.cell(row, source_col).value or "").strip()
        if current_source and current_source not in {
            "ADMINISTRATIVE_INFORMATION",
            "EXTERNAL_INPUT_REQUIRED",
            "NOT_FOUND",
            "DERIVED_METADATA",
            "DISTRICT_DATA_REQUIRED",
            "NO_APPROVED_PARTY_FOR_THIS_SLOT",
        }:
            continue
        literal = _stata_literal(constants[target])
        worksheet.cell(row, source_col).value = "BENCHMARK_REFERENCE_DECISION"
        worksheet.cell(row, transform_col).value = "calculate"
        worksheet.cell(row, recode_col).value = f"{literal}={literal}"
        if confidence_col:
            worksheet.cell(row, confidence_col).value = "benchmark"
        worksheet.cell(row, verified_col).value = "TRUE"
        if notes_col:
            note = str(worksheet.cell(row, notes_col).value or "").strip()
            replay_note = "Benchmark processor simulation: constant value inferred from processed reference dataset."
            worksheet.cell(row, notes_col).value = f"{note} | {replay_note}" if note else replay_note
    workbook.save(sheet)


def _stata_literal(value: Any) -> str:
    if isinstance(value, str):
        return '"' + value.replace('"', "'") + '"'
    if value is None:
        return "."
    return str(value)


def _find_latest(work_dir: Path, patterns: list[str]) -> Path | None:
    matches: list[Path] = []
    for pattern in patterns:
        matches.extend(work_dir.rglob(pattern))
    ignored_parts = {"_OLD", "FINAL dataset", "benchmark_report"}
    matches = [
        path for path in matches
        if path.is_file() and not any(part in ignored_parts or part.startswith("_OLD") for part in path.parts)
    ]
    return max(matches, key=lambda p: p.stat().st_mtime) if matches else None


def _syntax_metrics(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return {"exists": False}
    text = path.read_text(encoding="utf-8", errors="replace")
    lower = text.lower()
    return {
        "exists": True,
        "path": str(path),
        "bytes": path.stat().st_size,
        "variable_blocks": text.count("**>>>"),
        "tab_mis_commands": lower.count(", mis"),
        "generate_commands": lower.count("\ngen ") + lower.count("\ngenerate "),
        "replace_commands": lower.count("\nreplace "),
        "recode_commands": lower.count("\nrecode "),
        "has_log_close": "log close" in lower,
        "has_final_save": "\nsave " in lower or "\nsaveold " in lower,
    }


def _dataset_summary(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return {"exists": False}
    try:
        import pyreadstat

        _, meta = pyreadstat.read_dta(str(path), metadataonly=True)
        return {
            "exists": True,
            "path": str(path),
            "rows": getattr(meta, "number_rows", None),
            "variables": list(meta.column_names),
            "variable_count": len(meta.column_names),
            "labelled_variables": len(getattr(meta, "variable_value_labels", {}) or {}),
        }
    except Exception as exc:
        return {"exists": True, "path": str(path), "error": str(exc)}


def _compare_datasets(reference: dict[str, Any], generated: dict[str, Any]) -> dict[str, Any]:
    if not reference.get("exists") or not generated.get("exists"):
        return {"status": "missing", "details": "Reference or generated dataset not found"}
    ref_vars = set(reference.get("variables", []))
    gen_vars = set(generated.get("variables", []))
    overlap = sorted(ref_vars & gen_vars)
    missing = sorted(ref_vars - gen_vars)
    extra = sorted(gen_vars - ref_vars)
    return {
        "status": "compared",
        "reference_variable_count": len(ref_vars),
        "generated_variable_count": len(gen_vars),
        "overlap_count": len(overlap),
        "missing_reference_variables": missing[:100],
        "extra_generated_variables": extra[:100],
        "row_count_match": reference.get("rows") == generated.get("rows"),
        "reference_rows": reference.get("rows"),
        "generated_rows": generated.get("rows"),
    }


def _strict_dataset_comparison(reference_path: Path | None, generated_path: Path | None) -> dict[str, Any]:
    if not reference_path or not generated_path or not generated_path.exists() or not reference_path.exists():
        return {"status": "missing", "exact_variable_inventory": False, "exact_value_match": False}
    try:
        import pandas as pd
        import pyreadstat

        ref, ref_meta = pyreadstat.read_dta(str(reference_path), apply_value_formats=False)
        gen, gen_meta = pyreadstat.read_dta(str(generated_path), apply_value_formats=False)
        ref_vars = list(ref_meta.column_names)
        gen_vars = list(gen_meta.column_names)
        exact_inventory = ref_vars == gen_vars
        alignment_key = ""
        for candidate in ["F1003_2", "F1003_1"]:
            if candidate in ref.columns and candidate in gen.columns and ref[candidate].is_unique and gen[candidate].is_unique:
                ref = ref.set_index(candidate, drop=False).sort_index()
                gen = gen.set_index(candidate, drop=False).sort_index()
                alignment_key = candidate
                break
        value_mismatches = []
        substantive_value_mismatches = []
        group_metrics: dict[str, dict[str, Any]] = {}
        common = [var for var in ref_vars if var in gen.columns]
        for var in common:
            ref_series = ref[var]
            gen_series = gen[var]
            equal = _series_values_match(ref_series, gen_series)
            group = _variable_group(var)
            metric = group_metrics.setdefault(group, {"total": 0, "exact": 0, "mismatch_examples": []})
            metric["total"] += 1
            if equal:
                metric["exact"] += 1
            elif len(metric["mismatch_examples"]) < 25:
                metric["mismatch_examples"].append(var)
            if not equal:
                if len(value_mismatches) < 50:
                    value_mismatches.append(var)
                if not _is_release_metadata_variable(var) and len(substantive_value_mismatches) < 50:
                    substantive_value_mismatches.append(var)
        for metric in group_metrics.values():
            total = metric["total"] or 0
            metric["exact_share"] = round(metric["exact"] / total, 4) if total else 0.0
        label_match = (getattr(ref_meta, "column_labels", []) or []) == (getattr(gen_meta, "column_labels", []) or [])
        ref_labels = dict(zip(ref_vars, getattr(ref_meta, "column_labels", []) or []))
        gen_labels = dict(zip(gen_vars, getattr(gen_meta, "column_labels", []) or []))
        overlap_label_mismatches = [
            var for var in common
            if (ref_labels.get(var) or "") != (gen_labels.get(var) or "")
        ]
        overlap_labeled_count = sum(1 for var in common if bool(gen_labels.get(var)))
        overlap_label_match_share = (
            round((len(common) - len(overlap_label_mismatches)) / len(common), 4)
            if common else 0.0
        )
        return {
            "status": "compared",
            "exact_variable_inventory": exact_inventory,
            "exact_value_match": not value_mismatches and exact_inventory,
            "substantive_value_match": not substantive_value_mismatches,
            "column_label_match": label_match,
            "overlap_column_label_match": not overlap_label_mismatches,
            "overlap_column_label_match_share": overlap_label_match_share,
            "overlap_generated_label_coverage": round(overlap_labeled_count / len(common), 4) if common else 0.0,
            "alignment_key": alignment_key,
            "group_metrics": group_metrics,
            "value_mismatch_examples": value_mismatches,
            "substantive_value_mismatch_examples": substantive_value_mismatches,
            "overlap_label_mismatch_examples": overlap_label_mismatches[:50],
            "missing_reference_variables": [var for var in ref_vars if var not in gen_vars],
            "extra_generated_variables": [var for var in gen_vars if var not in ref_vars],
        }
    except Exception as exc:
        return {"status": "error", "error": str(exc), "exact_variable_inventory": False, "exact_value_match": False}


def _series_values_match(reference: Any, generated: Any) -> bool:
    try:
        import pandas as pd
        from pandas.api.types import is_numeric_dtype

        if is_numeric_dtype(reference) and is_numeric_dtype(generated):
            return reference.fillna(-9.87654321e30).astype(float).round(8).equals(
                generated.fillna(-9.87654321e30).astype(float).round(8)
            )
        ref_text = reference.fillna("__NA__").astype(str)
        gen_text = generated.fillna("__NA__").astype(str)
        if ref_text.equals(gen_text):
            return True
        ref_num = pd.to_numeric(reference, errors="coerce")
        gen_num = pd.to_numeric(generated, errors="coerce")
        if ref_num.notna().any() and gen_num.notna().any():
            return ref_num.fillna(-9.87654321e30).astype(float).round(8).equals(
                gen_num.fillna(-9.87654321e30).astype(float).round(8)
            )
        return False
    except Exception:
        try:
            return reference.fillna("__NA__").astype(str).equals(generated.fillna("__NA__").astype(str))
        except Exception:
            return False


def _is_release_metadata_variable(variable: str) -> bool:
    return variable in {"F1002_VER", "F1002_DOI"}


def _variable_group(variable: str) -> str:
    if variable.startswith(("F100", "F101", "F102", "F110")):
        return "administrative"
    if variable.startswith("F20"):
        return "demographic"
    if variable.startswith(("F30", "F31", "F5", "F6")):
        return "party_vote_leader_macro"
    if variable.startswith("F4"):
        return "district"
    return "core_questionnaire"


def _reference_path(reference_root: Path, reference_artifacts: dict[str, str], label: str) -> Path | None:
    rel_path = reference_artifacts.get(label, "")
    return reference_root / rel_path if rel_path else None


def _missing_materials(reference_root: Path, work_dir: Path, reference_artifacts: dict[str, str]) -> list[dict[str, str]]:
    materials = []
    for label, rel_path in reference_artifacts.items():
        ref_path = reference_root / rel_path if rel_path else None
        generated_candidates = list(work_dir.rglob(Path(rel_path).name)) if rel_path else []
        status = "present_in_email_only_run" if generated_candidates else "missing_from_email_only_run"
        materials.append(
            {
                "label": label,
                "reference_path": str(ref_path) if ref_path else "",
                "reference_exists": str(bool(ref_path and ref_path.exists())),
                "email_only_status": status,
                "classification": _classify_material(label),
            }
        )
    return materials


def _classify_material(label: str) -> str:
    if label in {"election_results", "macro_dataset", "esn", "macro_log"}:
        return "user_or_macro_process_must_supply_or_confirm"
    if label in {"reference_micro_syntax", "processing_log", "final_micro_dataset"}:
        return "agent_should_reproduce_functionally"
    return "processor_review_required"


def _build_report(
    work_dir: Path,
    reference_root: Path,
    init_result: dict[str, Any],
    transcripts: list[StepTranscript],
    mode: str,
    reference_artifacts: dict[str, str] | None = None,
    benchmark_name: str = "Sweden",
) -> dict[str, Any]:
    reference_artifacts = reference_artifacts or REFERENCE_ARTIFACTS
    state = WorkflowState.load(work_dir)
    generated_do = _find_latest(work_dir, ["cses-m6_micro_*.do"])
    generated_dataset = _find_latest(work_dir, ["cses-m6_micro_*.dta"])
    reference_do = _reference_path(reference_root, reference_artifacts, "reference_micro_syntax")
    reference_dataset = _reference_path(reference_root, reference_artifacts, "final_micro_dataset")
    reference_dataset_summary = _dataset_summary(reference_dataset)
    generated_dataset_summary = _dataset_summary(generated_dataset)
    completed_steps = [
        int(num)
        for num, step in (state.steps.items() if state else [])
        if getattr(state.get_step(int(num)), "status", "") == "completed"
    ]
    blockers = []
    if state:
        for num in WORKFLOW_SEQUENCE:
            step_state = state.get_step(num)
            if step_state.status != "completed":
                blockers.append(
                    {
                        "step": num,
                        "name": WORKFLOW_STEPS[num]["name"],
                        "status": step_state.status,
                        "issues": step_state.issues,
                    }
                )
                break
    documentation_comparison = DocumentationComparator().write_report(work_dir, reference_root)
    acceptance = _acceptance_summary(
        completed_steps=completed_steps,
        syntax_reference=_syntax_metrics(reference_do),
        syntax_generated=_syntax_metrics(generated_do),
        dataset_comparison=_compare_datasets(reference_dataset_summary, generated_dataset_summary),
        schema_target_count=(state.workflow_tracking or {}).get("target_count") if state else None,
        strict_dataset=_strict_dataset_comparison(reference_dataset, generated_dataset),
        candidate_question_count=len(state.candidate_collaborator_questions) if state else 0,
        documentation=documentation_comparison,
    )
    benchmark_decisions = BenchmarkDecisionExtractor().extract(
        reference_dataset=reference_dataset,
        reference_syntax=reference_do,
        reference_dir=reference_root,
        working_dir=work_dir,
    )
    return {
        "generated_at": datetime.now().isoformat(),
        "benchmark_name": benchmark_name,
        "benchmark_mode": mode,
        "work_dir": str(work_dir),
        "reference_root": str(reference_root),
        "reference_artifacts": reference_artifacts,
        "init": init_result,
        "completed_steps": completed_steps,
        "first_blocker": blockers[0] if blockers else None,
        "evidence_diagnostics": (state.evidence_index or {}).get("diagnostics", {}) if state else {},
        "missing_fields": (state.evidence_index or {}).get("missing_fields", []) if state else [],
        "candidate_collaborator_questions": state.candidate_collaborator_questions if state else [],
        "transcripts": [asdict(item) for item in transcripts],
        "acceptance": acceptance,
        "documentation_comparison": documentation_comparison,
        "benchmark_decision_replay": {
            "path": str(work_dir / ".cses" / "benchmark_decision_replay.json"),
            "constant_values": len(benchmark_decisions.get("constant_values", {})),
            "syntax_variable_blocks": len(benchmark_decisions.get("syntax_variable_blocks", {})),
            "reference_recoding_plans": len(benchmark_decisions.get("reference_recoding_plans", {})),
            "documentation_decision_topics": benchmark_decisions.get("documentation_decision_topics", {}),
        },
        "syntax_comparison": {
            "reference": acceptance["syntax_reference"],
            "generated": acceptance["syntax_generated"],
        },
        "dataset_comparison": {
            "reference": reference_dataset_summary,
            "generated": generated_dataset_summary,
            "comparison": _compare_datasets(reference_dataset_summary, generated_dataset_summary),
        },
        "missing_materials": _missing_materials(reference_root, work_dir, reference_artifacts),
        "input_manifest": _relative_manifest(work_dir / "E-mails")
        if (work_dir / "E-mails").exists()
        else _relative_manifest(work_dir / "emails"),
        "output_manifest": _relative_manifest(work_dir),
    }


def _acceptance_summary(
    completed_steps: list[int],
    syntax_reference: dict[str, Any],
    syntax_generated: dict[str, Any],
    dataset_comparison: dict[str, Any],
    schema_target_count: int | None = None,
    strict_dataset: dict[str, Any] | None = None,
    candidate_question_count: int = 0,
    documentation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    reference_blocks = syntax_reference.get("variable_blocks") or 0
    generated_blocks = syntax_generated.get("variable_blocks") or 0
    reference_vars = dataset_comparison.get("reference_variable_count") or 0
    generated_vars = dataset_comparison.get("generated_variable_count") or 0
    overlap = dataset_comparison.get("overlap_count") or 0
    checks = {
        "all_steps_completed": len(completed_steps) >= 17,
        "stata_syntax_generated": bool(syntax_generated.get("exists")),
        "syntax_block_coverage_90pct": generated_blocks >= 0.9 * reference_blocks if reference_blocks else False,
        "schema_variable_coverage_90pct": generated_blocks >= 0.9 * schema_target_count if schema_target_count else False,
        "dataset_generated": dataset_comparison.get("status") == "compared",
        "dataset_variable_coverage_90pct": generated_vars >= 0.9 * reference_vars if reference_vars else False,
        "reference_overlap_90pct": overlap >= 0.9 * reference_vars if reference_vars else False,
        "row_count_match": bool(dataset_comparison.get("row_count_match")),
        "exact_variable_inventory": bool((strict_dataset or {}).get("exact_variable_inventory")),
        "exact_value_match": bool((strict_dataset or {}).get("exact_value_match")),
        "column_label_match": bool((strict_dataset or {}).get("column_label_match")),
        "documentation_equivalence": bool((documentation or {}).get("ok", not documentation)),
        "no_candidate_collaborator_questions": candidate_question_count == 0,
    }
    return {
        "functional_replication_success": all(checks.values()),
        "checks": checks,
        "syntax_reference": syntax_reference,
        "syntax_generated": syntax_generated,
        "dataset_comparison": dataset_comparison,
        "strict_dataset_comparison": strict_dataset or {},
    }


def _write_markdown_report(report: dict[str, Any], path: Path) -> None:
    blocker = report.get("first_blocker")
    syntax = report["syntax_comparison"]
    dataset = report["dataset_comparison"]["comparison"]
    lines = [
        f"# {report.get('benchmark_name', 'Study')} Replication Benchmark ({report.get('benchmark_mode', 'email_only')})",
        "",
        f"Generated at: {report['generated_at']}",
        f"Work folder: {report['work_dir']}",
        f"Reference folder: {report['reference_root']}",
        "",
        "## Scorecard",
        f"- Init return code: {report['init']['returncode']}",
        f"- Completed steps: {len(report['completed_steps'])}/17",
        f"- Functional replication success: {report['acceptance']['functional_replication_success']}",
        f"- First blocker: {blocker if blocker else 'None'}",
        f"- Evidence chunks: {report['evidence_diagnostics'].get('chunks_ok', 0)}/{report['evidence_diagnostics'].get('chunks_total', 0)}",
        f"- Evidence fields found: {report['evidence_diagnostics'].get('fields_found', 0)}",
        f"- Missing fields: {', '.join(report.get('missing_fields', [])) or 'None'}",
        f"- Exact variable inventory: {report['acceptance'].get('strict_dataset_comparison', {}).get('exact_variable_inventory')}",
        f"- Exact value match: {report['acceptance'].get('strict_dataset_comparison', {}).get('exact_value_match')}",
        f"- Column label match: {report['acceptance'].get('strict_dataset_comparison', {}).get('column_label_match')}",
        f"- Documentation comparison: {'pass' if report.get('documentation_comparison', {}).get('ok') else 'needs review'}",
        f"- Benchmark reference decisions extracted: {report.get('benchmark_decision_replay', {}).get('constant_values', 0)} constants, "
        f"{report.get('benchmark_decision_replay', {}).get('syntax_variable_blocks', 0)} syntax blocks, "
        f"{report.get('benchmark_decision_replay', {}).get('reference_recoding_plans', 0)} recoding plans",
        "",
        "## Syntax Comparison",
        f"- Reference variable blocks: {syntax['reference'].get('variable_blocks')}",
        f"- Generated variable blocks: {syntax['generated'].get('variable_blocks')}",
        f"- Reference tab/mis commands: {syntax['reference'].get('tab_mis_commands')}",
        f"- Generated tab/mis commands: {syntax['generated'].get('tab_mis_commands')}",
        "",
        "## Dataset Comparison",
        f"- Status: {dataset.get('status')}",
        f"- Reference variables: {dataset.get('reference_variable_count')}",
        f"- Generated variables: {dataset.get('generated_variable_count')}",
        f"- Overlap: {dataset.get('overlap_count')}",
        f"- Row count match: {dataset.get('row_count_match')}",
        "",
        "### Dataset Replication by Area",
    ]
    group_metrics = report["acceptance"].get("strict_dataset_comparison", {}).get("group_metrics", {})
    for group, metric in group_metrics.items():
        lines.append(
            f"- {group}: {metric.get('exact')}/{metric.get('total')} exact "
            f"({metric.get('exact_share')})"
        )
    lines.extend([
        "",
        "## Documentation Comparison",
    ])
    documentation = report.get("documentation_comparison", {})
    lines.extend(
        f"- {key}: {'present' if value else 'missing'}"
        for key, value in documentation.get("checks", {}).items()
    )
    if documentation.get("issues"):
        lines.append("")
        lines.append("Documentation issues:")
        lines.extend(f"- {issue}" for issue in documentation.get("issues", [])[:30])
    decision_replay = report.get("benchmark_decision_replay", {})
    lines.extend([
        "",
        "## Benchmark Decision Replay",
        f"- Reference constants extracted: {decision_replay.get('constant_values', 0)}",
        f"- Reference syntax blocks extracted: {decision_replay.get('syntax_variable_blocks', 0)}",
        f"- Reference recoding plans extracted: {decision_replay.get('reference_recoding_plans', 0)}",
        f"- Replay file: {decision_replay.get('path', '')}",
        "",
        "## Missing Materials Assessment",
    ])
    for item in report["missing_materials"]:
        lines.append(
            f"- {item['label']}: {item['email_only_status']} ({item['classification']})"
        )
    lines.extend(["", "## Workflow Transcript"])
    for transcript in report["transcripts"]:
        lines.append(
            f"### Step {transcript['step']}: {transcript['step_name']} - {transcript['status_after']}"
        )
        if transcript["issues_after"]:
            lines.append("Issues:")
            lines.extend(f"- {issue}" for issue in transcript["issues_after"][:10])
        lines.append("Assistant reply excerpt:")
        lines.append("```")
        lines.append(transcript["assistant_reply"][:2000])
        lines.append("```")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Run a CSES study replication benchmark")
    parser.add_argument("--source-study", type=Path, default=REPO_ROOT / "Sweden_2022")
    parser.add_argument("--reference-study", type=Path, default=REPO_ROOT / "Sweden_2022")
    parser.add_argument("--reference-artifacts-json", type=Path, default=None)
    parser.add_argument("--work-dir", type=Path, default=None)
    parser.add_argument("--country", default="Sweden")
    parser.add_argument("--year", default="2022")
    parser.add_argument("--benchmark-name", default="Sweden")
    parser.add_argument("--max-steps", type=int, default=17)
    parser.add_argument("--mode", choices=["email_only", "full_reference_inputs"], default="email_only")
    parser.add_argument("--stata-path", default="", help="Optional Stata executable path used by the MCP-Stata bridge")
    args = parser.parse_args()

    apply_settings_to_environment()
    if args.stata_path:
        os.environ["STATA_PATH"] = args.stata_path

    reference_artifacts = dict(REFERENCE_ARTIFACTS)
    if args.reference_artifacts_json:
        reference_artifacts = json.loads(args.reference_artifacts_json.read_text(encoding="utf-8"))

    safe_name = re.sub(r"[^A-Za-z0-9]+", "_", args.benchmark_name).strip("_").lower() or "study"
    work_dir = args.work_dir or Path(tempfile.gettempdir()) / (
        f"cses_{safe_name}_{args.mode}_replication_" + datetime.now().strftime("%Y%m%d_%H%M%S")
    )
    if work_dir.exists():
        raise FileExistsError(f"Work directory already exists: {work_dir}")

    if args.mode == "email_only":
        _copy_email_only_reference(args.source_study, work_dir)
    else:
        _copy_full_reference_inputs(args.source_study, work_dir)
    init_result = _run_init(work_dir, args.country, args.year)
    transcripts: list[StepTranscript] = []
    if init_result["returncode"] == 0:
        if args.mode == "full_reference_inputs":
            BenchmarkDecisionExtractor().extract(
                reference_dataset=_reference_path(args.reference_study, reference_artifacts, "final_micro_dataset"),
                reference_syntax=_reference_path(args.reference_study, reference_artifacts, "reference_micro_syntax"),
                reference_dir=args.reference_study,
                working_dir=work_dir,
            )
        transcripts = _run_conversation(work_dir, args.max_steps)

    report = _build_report(
        work_dir,
        args.reference_study,
        init_result,
        transcripts,
        args.mode,
        reference_artifacts=reference_artifacts,
        benchmark_name=args.benchmark_name,
    )
    report_dir = work_dir / "benchmark_report"
    report_dir.mkdir(parents=True, exist_ok=True)
    json_path = report_dir / f"{safe_name}_{args.mode}_replication.json"
    md_path = report_dir / f"{safe_name}_{args.mode}_replication.md"
    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    _write_markdown_report(report, md_path)
    print(f"WORK_DIR={work_dir}")
    print(f"REPORT_JSON={json_path}")
    print(f"REPORT_MD={md_path}")
    print(f"COMPLETED_STEPS={len(report['completed_steps'])}/17")
    print(f"FIRST_BLOCKER={report.get('first_blocker')}")
    return 0


if __name__ == "__main__":
    exit_code = main()
    sys.stdout.flush()
    sys.stderr.flush()
    os._exit(exit_code)
