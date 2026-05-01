"""Election-results parsing and party-order agreement support.

Party order is a joint micro/macro coding decision. This module only proposes
and validates the order from generic evidence; it does not silently approve it.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


PARTY_LETTERS = tuple("ABCDEFGHI")


@dataclass
class ElectionResultParty:
    party_label: str = ""
    code_letter: str = ""
    numeric_code: str = ""
    party_name: str = ""
    votes: float | None = None
    vote_share: float | None = None
    seats: float | None = None
    seat_share: float | None = None
    row_number: int = 0
    is_others: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class ElectionResultTable:
    source_file: str
    sheet_name: str = ""
    title: str = ""
    election_context: str = "unknown"
    parties: list[ElectionResultParty] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["parties"] = [party.to_dict() for party in self.parties]
        return data


@dataclass
class PartyOrderProposal:
    status: str
    selected_context: str = "unknown"
    ordering_rule: str = ""
    ordering_basis: str = ""
    source_file: str = ""
    source_sheet: str = ""
    proposed_parties: list[ElectionResultParty] = field(default_factory=list)
    affected_micro_variables: list[str] = field(default_factory=list)
    affected_macro_variables: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    processor_decision_needed: str = ""

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["proposed_parties"] = [party.to_dict() for party in self.proposed_parties]
        return data


class ElectionResultsWorkbookParser:
    """Parse standardized CSES election-results workbooks without country logic."""

    HEADER_ALIASES = {
        "party_label": ("party label", "party", "party slot"),
        "numeric_code": ("numeric", "numeric code", "party code", "cses code"),
        "party_name": ("party name", "name"),
        "votes": ("votes", "vote"),
        "vote_share": ("% of vote", "percent vote", "vote share", "% vote"),
        "seats": ("seats", "seat"),
        "seat_share": ("% of seats", "percent seats", "seat share", "% seats"),
    }

    def discover_files(self, working_dir: Path) -> list[Path]:
        roots = [
            Path(working_dir),
            Path(working_dir) / "Election Results",
            Path(working_dir) / "macro",
            Path(working_dir) / "E-mails",
            Path(working_dir) / "emails",
        ]
        patterns = (
            "*Election*Results*.xlsx",
            "*election*results*.xlsx",
            "*election_result*.xlsx",
            "*results*.xlsx",
            "*Election*Results*.csv",
            "*election*results*.csv",
            "*results*.csv",
        )
        found: dict[str, Path] = {}
        for root in roots:
            if not root.exists():
                continue
            for pattern in patterns:
                for path in root.rglob(pattern):
                    if path.is_file() and not path.name.startswith("~$"):
                        found[str(path.resolve()).lower()] = path
        return sorted(found.values(), key=lambda item: str(item).lower())

    def parse_file(self, path: Path) -> list[ElectionResultTable]:
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            return self._parse_xlsx(path)
        if suffix == ".csv":
            return self._parse_csv(path)
        return []

    def parse_directory(self, working_dir: Path) -> list[ElectionResultTable]:
        tables: list[ElectionResultTable] = []
        for path in self.discover_files(working_dir):
            tables.extend(self.parse_file(path))
        return tables

    def _parse_xlsx(self, path: Path) -> list[ElectionResultTable]:
        try:
            import openpyxl
        except ImportError:
            return []
        tables: list[ElectionResultTable] = []
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            tables.extend(self._parse_rows(rows, path, ws.title))
        wb.close()
        return tables

    def _parse_csv(self, path: Path) -> list[ElectionResultTable]:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return self._parse_rows(rows, path, path.stem)

    def _parse_rows(self, rows: list[list[Any]], path: Path, sheet_name: str) -> list[ElectionResultTable]:
        tables: list[ElectionResultTable] = []
        for row_index, row in enumerate(rows):
            header_map = self._header_map(row)
            if not {"party_name", "vote_share"}.issubset(header_map):
                continue
            title = self._nearest_title(rows, row_index)
            parties: list[ElectionResultParty] = []
            for data_index in range(row_index + 1, len(rows)):
                data_row = rows[data_index]
                if self._is_blank(data_row):
                    if parties:
                        break
                    continue
                party = self._party_from_row(data_row, header_map, data_index + 1)
                if not party:
                    if parties:
                        break
                    continue
                parties.append(party)
            if parties:
                tables.append(
                    ElectionResultTable(
                        source_file=str(path),
                        sheet_name=sheet_name,
                        title=title,
                        election_context=self._classify_context(" ".join([title, sheet_name])),
                        parties=parties,
                        warnings=self._table_warnings(parties),
                    )
                )
        return tables

    def _header_map(self, row: list[Any]) -> dict[str, int]:
        mapped: dict[str, int] = {}
        for index, cell in enumerate(row):
            label = _norm_header(cell)
            if not label:
                continue
            for key, aliases in self.HEADER_ALIASES.items():
                if any(label == alias or alias in label for alias in aliases):
                    mapped.setdefault(key, index)
        return mapped

    def _party_from_row(self, row: list[Any], header_map: dict[str, int], row_number: int) -> ElectionResultParty | None:
        name = _cell(row, header_map.get("party_name"))
        label = _cell(row, header_map.get("party_label"))
        numeric = _cell(row, header_map.get("numeric_code"))
        vote_share = _number(_cell(row, header_map.get("vote_share")))
        votes = _number(_cell(row, header_map.get("votes")))
        seats = _number(_cell(row, header_map.get("seats")))
        seat_share = _number(_cell(row, header_map.get("seat_share")))
        if not any([name, label, numeric, vote_share is not None, votes is not None]):
            return None
        is_others = bool(re.search(r"\bothers?\b", name, re.IGNORECASE))
        return ElectionResultParty(
            party_label=label,
            code_letter=_party_letter(label),
            numeric_code=_clean_numeric_code(numeric),
            party_name=name,
            votes=votes,
            vote_share=vote_share,
            seats=seats,
            seat_share=seat_share,
            row_number=row_number,
            is_others=is_others,
        )

    def _nearest_title(self, rows: list[list[Any]], header_row: int) -> str:
        for index in range(header_row - 1, max(-1, header_row - 6), -1):
            cells = [str(cell).strip() for cell in rows[index] if str(cell or "").strip()]
            if cells:
                return " | ".join(cells[:3])
        return ""

    def _classify_context(self, text: str) -> str:
        normalized = text.casefold()
        if any(token in normalized for token in ("president", "presidential")):
            if "second" in normalized or "round 2" in normalized:
                return "presidential_round_2"
            return "presidential_round_1"
        if any(token in normalized for token in ("upper house", "senate", "upper chamber")):
            return "upper_house"
        if any(token in normalized for token in ("lower house", "parliament", "assembly", "house")):
            return "lower_house"
        return "unknown"

    def _table_warnings(self, parties: list[ElectionResultParty]) -> list[str]:
        warnings: list[str] = []
        labeled = [party for party in parties if party.code_letter]
        if not labeled:
            warnings.append("No Party A-I labels were found in this table.")
        if any(not party.numeric_code for party in labeled):
            warnings.append("At least one Party A-I row is missing a numeric party code.")
        if any(party.vote_share is None for party in labeled):
            warnings.append("At least one Party A-I row is missing vote share.")
        return warnings

    def _is_blank(self, row: list[Any]) -> bool:
        return not any(str(cell or "").strip() for cell in row)


class PartyOrderingRulesEngine:
    """Propose party order and agreement requirements from parsed tables."""

    def __init__(self, rules_path: Path | None = None):
        self.rules_path = rules_path or Path(__file__).resolve().parents[2] / "cses_wiki" / "patterns" / "party_ordering_rules.json"
        self.rules = json.loads(self.rules_path.read_text(encoding="utf-8")) if self.rules_path.exists() else {}

    def propose(
        self,
        tables: list[ElectionResultTable],
        source_variables: list[str] | None = None,
        context_hint: str = "",
    ) -> PartyOrderProposal:
        if not tables:
            return PartyOrderProposal(
                status="missing_election_results",
                warnings=["No standardized election-results table was found."],
                affected_micro_variables=self._affected_micro_variables(source_variables or []),
                affected_macro_variables=_party_macro_variables(),
                processor_decision_needed="Provide the standardized election-results workbook or confirm an approved public-source lookup.",
            )
        table = self._select_table(tables)
        selected_context = table.election_context
        if selected_context == "unknown" and context_hint:
            selected_context = context_hint
        labeled = [party for party in table.parties if party.code_letter in PARTY_LETTERS and not party.is_others]
        if not labeled:
            labeled = self._assign_letters_from_vote_share(table.parties)
        computed = sorted(
            [party for party in labeled if party.vote_share is not None],
            key=lambda party: (-float(party.vote_share or 0), party.party_name.casefold()),
        )
        proposed = computed if computed else labeled
        warnings = [*table.warnings]
        if labeled and computed:
            workbook_order = [party.code_letter for party in labeled]
            computed_order = [party.code_letter for party in computed]
            if workbook_order != computed_order:
                warnings.append("Workbook Party A-I order differs from descending vote-share order.")
        if len(proposed) > len(PARTY_LETTERS):
            warnings.append("More than nine Party A-I candidates were found; processor must select the CSES party table.")
            proposed = proposed[: len(PARTY_LETTERS)]
        return PartyOrderProposal(
            status="agreement_needed",
            selected_context=selected_context,
            ordering_rule=self._rule_id(selected_context),
            ordering_basis=self._basis_text(selected_context),
            source_file=table.source_file,
            source_sheet=table.sheet_name,
            proposed_parties=proposed,
            affected_micro_variables=self._affected_micro_variables(source_variables or []),
            affected_macro_variables=_party_macro_variables(),
            warnings=warnings,
            processor_decision_needed="Micro processor and macro coder must approve this party order before party, vote-choice, leader, and macro-party coding proceeds.",
        )

    def write_review(self, working_dir: Path, proposal: PartyOrderProposal) -> tuple[Path, Path]:
        cses_dir = Path(working_dir) / ".cses"
        cses_dir.mkdir(parents=True, exist_ok=True)
        review_path = cses_dir / "party_order_review.json"
        decision_path = cses_dir / "party_order_decision.json"
        payload = {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "proposal": proposal.to_dict(),
            "approval": {
                "micro_processor_approved": False,
                "macro_coder_approved": False,
                "locked": False,
                "override_reason": "",
            },
        }
        review_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        if not decision_path.exists():
            decision_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        return review_path, decision_path

    def load_decision(self, working_dir: Path) -> dict[str, Any]:
        path = Path(working_dir) / ".cses" / "party_order_decision.json"
        if not path.exists():
            return {}
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def is_approved(self, working_dir: Path) -> bool:
        decision = self.load_decision(working_dir)
        approval = decision.get("approval", {}) if isinstance(decision, dict) else {}
        return bool(
            approval.get("locked")
            and approval.get("micro_processor_approved")
            and approval.get("macro_coder_approved")
        )

    def _select_table(self, tables: list[ElectionResultTable]) -> ElectionResultTable:
        priority = {"lower_house": 0, "presidential_round_1": 1, "upper_house": 2, "presidential_round_2": 3, "unknown": 4}
        return sorted(
            tables,
            key=lambda table: (
                priority.get(table.election_context, 9),
                -len([party for party in table.parties if party.code_letter]),
                str(table.source_file).casefold(),
            ),
        )[0]

    def _rule_id(self, context: str) -> str:
        return {
            "lower_house": "lower_house_vote_share",
            "upper_house": "upper_house_vote_share_when_applicable",
            "presidential_round_1": "presidential_first_round_vote_share",
            "presidential_round_2": "presidential_second_round_context_only",
        }.get(context, "processor_review_required")

    def _basis_text(self, context: str) -> str:
        rules = self.rules.get("rules", {}) if isinstance(self.rules, dict) else {}
        rule = rules.get(self._rule_id(context), {})
        return str(rule.get("ordering_basis") or rule.get("description") or "Processor review required.")

    def _assign_letters_from_vote_share(self, parties: list[ElectionResultParty]) -> list[ElectionResultParty]:
        eligible = [
            party for party in parties
            if not party.is_others and party.party_name and party.vote_share is not None
        ]
        ordered = sorted(eligible, key=lambda party: (-float(party.vote_share or 0), party.party_name.casefold()))
        assigned: list[ElectionResultParty] = []
        for index, party in enumerate(ordered[: len(PARTY_LETTERS)]):
            assigned_party = ElectionResultParty(**party.to_dict())
            assigned_party.code_letter = PARTY_LETTERS[index]
            assigned_party.party_label = f"PARTY {PARTY_LETTERS[index]}"
            assigned.append(assigned_party)
        return assigned

    def _affected_micro_variables(self, source_variables: list[str]) -> list[str]:
        observed = [name for name in source_variables if _looks_party_source(name)]
        return sorted(set(_party_micro_variables() + observed))


def election_results_intake_summary(working_dir: Path) -> dict[str, Any]:
    parser = ElectionResultsWorkbookParser()
    files = parser.discover_files(working_dir)
    tables: list[ElectionResultTable] = []
    for path in files:
        tables.extend(parser.parse_file(path))
    return {
        "files": [str(path) for path in files],
        "tables": [table.to_dict() for table in tables],
        "election_context_hint": infer_election_context_from_macro_material(working_dir),
        "table_count": len(tables),
        "standardized": bool(tables),
    }


def write_election_results_intake(working_dir: Path, summary: dict[str, Any]) -> Path:
    path = Path(working_dir) / ".cses" / "election_results_intake.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"generated_at": datetime.now(timezone.utc).isoformat(), **summary}
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def infer_election_context_from_macro_material(working_dir: Path) -> str:
    """Infer broad election context from generic CSES macro variable evidence."""
    candidates: list[Path] = []
    roots = [Path(working_dir), Path(working_dir) / "macro"]
    for root in roots:
        if not root.exists():
            continue
        candidates.extend(path for path in root.rglob("*.xlsx") if path.is_file() and not path.name.startswith("~$"))
    for path in sorted(set(candidates), key=lambda item: str(item).lower()):
        context = _infer_context_from_workbook(path)
        if context:
            return context
    return ""


def party_agreement_summary(proposal: PartyOrderProposal) -> dict[str, int]:
    return {
        "party_count": len(proposal.proposed_parties),
        "warnings": len(proposal.warnings),
        "micro_variables_affected": len(proposal.affected_micro_variables),
        "macro_variables_affected": len(proposal.affected_macro_variables),
    }


def party_order_message(proposal: PartyOrderProposal) -> str:
    if proposal.status == "missing_election_results":
        return (
            "Party Order Agreement cannot be prepared yet.\n\n"
            "Election-results material needed:\n"
            "- Standardized election-results workbook with party names, vote shares, seats, and CSES party codes.\n\n"
            "Processor decision needed:\n"
            "- Provide the standardized workbook or confirm an approved public-source lookup."
        )
    lines = [
        "Party Order Agreement prepared.",
        "",
        f"Election context: {proposal.selected_context.replace('_', ' ')}",
        f"Ordering basis: {proposal.ordering_basis}",
        "",
        "Proposed party order:",
    ]
    for party in proposal.proposed_parties:
        vote = "" if party.vote_share is None else f", vote share {party.vote_share:g}"
        seats = "" if party.seats is None else f", seats {party.seats:g}"
        code = f", code {party.numeric_code}" if party.numeric_code else ""
        lines.append(f"- Party {party.code_letter}: {party.party_name}{code}{vote}{seats}")
    if proposal.warnings:
        lines.extend(["", "Needs review:"])
        lines.extend(f"- {warning}" for warning in proposal.warnings[:8])
    lines.extend([
        "",
        "Processor decision needed:",
        "- Micro processor and macro coder approve this order before party, vote-choice, leader, and macro-party coding.",
    ])
    return "\n".join(lines)


def is_party_order_dependent_variable(name: str, description: str = "", dependency_class: str = "", item_type: str = "") -> bool:
    if dependency_class == "macro_or_party_input":
        return True
    if item_type == "party_vote_item":
        return True
    if re.match(r"^F30(11|16|18|19|20|21|23)", name):
        return True
    return any(token in description.upper() for token in ("PARTY A", "PARTY B", "LEADER A", "VOTE CHOICE", "PARTY ID"))


def _party_micro_variables() -> list[str]:
    letters = PARTY_LETTERS
    variables = [
        "F3011_PR_1", "F3011_PR_2", "F3011_LH_PL", "F3011_LH_DC", "F3011_LH_PF",
        "F3011_UH_PL", "F3011_UH_DC_1", "F3011_UH_DC_2", "F3011_UH_PF",
        "F3011_OUTGOV", "F3011_IF_CSES", "F3016_PR_1", "F3016_PR_2",
        "F3016_LH_PL", "F3016_LH_DC", "F3016_UH_PL", "F3016_UH_DC_1", "F3016_UH_DC_2",
        "F3023_1", "F3023_2", "F3023_3", "F3023_4",
    ]
    for prefix in ("F3018", "F3019", "F3020", "F3021"):
        variables.extend(f"{prefix}_{letter}" for letter in letters)
    return variables


def _party_macro_variables() -> list[str]:
    letters = PARTY_LETTERS
    variables: list[str] = []
    for prefix in ("F5000", "F5000_L", "F5200", "F5201", "F5202", "F5203"):
        variables.extend(f"{prefix}_{letter}" for letter in letters)
    variables.extend(["F6000_PR_1", "F6000_PR_2", "F6000_LH_PL", "F6000_LH_DC"])
    return variables


def _looks_party_source(name: str) -> bool:
    return bool(re.search(r"(party|leader|vote|q10|q14|q16|q17|q18|q23)", str(name), re.IGNORECASE))


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _cell(row: list[Any], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: str) -> float | None:
    text = str(value or "").strip()
    if not text:
        return None
    text = text.replace("%", "").replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _party_letter(label: str) -> str:
    match = re.search(r"\b(?:party\s*)?([A-I])\b", str(label or ""), re.IGNORECASE)
    return match.group(1).upper() if match else ""


def _clean_numeric_code(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    numeric = _number(text)
    if numeric is not None and float(numeric).is_integer():
        return str(int(numeric))
    return text


def _infer_context_from_workbook(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""
    try:
        flags: dict[str, bool] = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(cell or "").strip() for cell in row]
                for index, cell in enumerate(cells):
                    if cell in {"F5001", "F5005", "F5006"}:
                        flags[cell] = _row_has_truthy_after(cells, index)
        if flags.get("F5001"):
            return "lower_house"
        if flags.get("F5005"):
            return "presidential_round_1"
        if flags.get("F5006"):
            return "presidential_round_2"
    finally:
        wb.close()
    return ""


def _row_has_truthy_after(cells: list[str], index: int) -> bool:
    for value in cells[index + 1:]:
        text = str(value or "").strip().casefold()
        if text in {"1", "yes", "true", "y"}:
            return True
    return False
