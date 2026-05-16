"""Schema-backed workflow tracking artifacts."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from src.standards.schema import SchemaRegistry, SchemaVariable
from src.workflow.state import WorkflowState


TRACKING_HEADERS = [
    "VARIABLES",
    "CSES code",
    "SOURCE VARIABLE(S)",
    "SOURCE EVIDENCE",
    "CONFIDENCE",
    "VERIFIED",
    "RECODING NOTE",
    "MISSING VALUE TREATMENT",
    "COLLABORATOR QUESTION",
    "REMARKS",
    "DEPENDENCY CLASS",
    "REQUIRED EVIDENCE",
    "SYNTAX STATUS",
    "STATA CHECK STATUS",
    "PROCESSOR DECISION",
    "WIKI PATTERN ID",
]


@dataclass
class TrackingRow:
    target_variable: str
    description: str
    dependency_class: str
    required_evidence: list[str]
    source_variables: str = "[not yet matched]"
    source_evidence: str = ""
    confidence: str = ""
    verified: str = "FALSE"
    recoding_note: str = ""
    missing_value_treatment: str = ""
    collaborator_question: str = ""
    remarks: str = ""
    syntax_status: str = "not_generated"
    stata_check_status: str = "not_run"
    processor_decision: str = ""
    wiki_pattern_id: str = ""

    @classmethod
    def from_schema(cls, item: SchemaVariable) -> "TrackingRow":
        return cls(
            target_variable=item.name,
            description=item.description,
            dependency_class=item.dependency_class,
            required_evidence=item.required_evidence,
            missing_value_treatment=item.missing_value_policy,
            wiki_pattern_id=item.syntax_pattern_id,
        )

    def to_dict(self) -> dict:
        return asdict(self)


class WorkflowTrackingModel:
    """Creates and persists the full CSES tracking contract."""

    def __init__(self, state: WorkflowState, wiki_root: Path | None = None):
        self.state = state
        self.registry = SchemaRegistry(wiki_root) if wiki_root else SchemaRegistry()

    def build_rows(self) -> list[TrackingRow]:
        rows = [TrackingRow.from_schema(item) for item in self.registry.variables]
        evidence = self.state.evidence_index or {}
        facts = evidence.get("facts_by_field", {})
        direct_variables = set(evidence.get("source_variables", []))
        mapping_lookup = {}
        for mapping in self.state.mappings or []:
            target = (
                mapping.get("target")
                or mapping.get("cses_target")
                or mapping.get("cses_variable")
                or mapping.get("target_variable")
            )
            if target:
                mapping_lookup[str(target)] = mapping
        for row in rows:
            mapping = mapping_lookup.get(row.target_variable)
            if mapping:
                source = (
                    mapping.get("source")
                    or mapping.get("source_variable")
                    or mapping.get("source_var")
                    or "NOT_FOUND"
                )
                row.source_variables = str(source)
                row.source_evidence = str(mapping.get("reasoning") or mapping.get("validation_reasoning") or "")
                row.confidence = str(mapping.get("confidence_level") or mapping.get("confidence") or "")
                row.verified = "TRUE" if row.confidence == "high" else "FALSE"
                row.recoding_note = "AI match; processor verification required before final release"
                if source in {"NOT_FOUND", "ERROR", ""}:
                    row.remarks = "No reliable source match found after ensemble review."
            elif row.target_variable in direct_variables:
                row.source_variables = row.target_variable
                row.source_evidence = "Direct CSES-named variable found in deposited data."
                row.confidence = "high"
                row.recoding_note = "direct copy candidate; processor verification required"
            elif row.dependency_class in {"macro_or_party_input", "district_input"}:
                row.source_variables = "EXTERNAL_INPUT_REQUIRED"
                row.confidence = "processor_review"
                row.remarks = "Requires macro/election/district material or documented processor decision."
            elif row.dependency_class == "derived_metadata":
                row.source_variables = "DERIVED_METADATA"
                row.confidence = "processor_review"
                row.remarks = "Derived from study metadata/design report and must be verified."
            if mapping and mapping.get("status") == "generated_from_administrative_information":
                row.source_variables = str(mapping.get("source_variable") or "ADMINISTRATIVE_INFORMATION")
                row.source_evidence = " | ".join(str(item) for item in mapping.get("evidence", [])[:3])
                row.confidence = str(mapping.get("confidence") or "processor_review")
                row.recoding_note = "Generated from reviewed administrative information"
                row.remarks = str(mapping.get("notes") or "")
            if facts:
                field_hits = _facts_for_row(row.target_variable, facts)
                if field_hits:
                    row.source_evidence = " | ".join(field_hits[:3])
                    row.confidence = row.confidence or "medium"
        return rows

    def to_state_payload(self) -> dict:
        rows = self.build_rows()
        return {
            "schema_version": 1,
            "target_count": len(rows),
            "rows": [row.to_dict() for row in rows],
        }

    def write_excel(self, output_path: Path) -> Path:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        rows = self.build_rows()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "deposited variables"

        fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        for col, header in enumerate(TRACKING_HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = fill

        for row_idx, row in enumerate(rows, 2):
            values = [
                row.description,
                row.target_variable,
                row.source_variables,
                row.source_evidence,
                row.confidence,
                row.verified,
                row.recoding_note,
                row.missing_value_treatment,
                row.collaborator_question,
                row.remarks,
                row.dependency_class,
                "; ".join(row.required_evidence),
                row.syntax_status,
                row.stata_check_status,
                row.processor_decision,
                row.wiki_pattern_id,
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row=row_idx, column=col, value=value)

        widths = [48, 16, 26, 48, 16, 10, 36, 30, 36, 38, 24, 40, 18, 20, 28, 26]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
        wb.save(output_path)
        return output_path


def _facts_for_row(target_variable: str, facts_by_field: dict) -> list[str]:
    terms = []
    if target_variable.startswith("F10") or target_variable.startswith("F11"):
        terms.extend(["sample_size", "sample_design", "weights", "fieldwork_dates", "mode"])
    if target_variable.startswith("F30"):
        terms.extend(["cses_item_coverage", "vote_choice_questions", "party_leader_questions"])
    if target_variable.startswith("F50") or target_variable.startswith("F60"):
        terms.extend(["election_date", "turnout", "party_evidence"])
    results = []
    for term in terms:
        for fact in facts_by_field.get(term, [])[:2]:
            source = Path(str(fact.get("source_file", ""))).name
            value = str(fact.get("value", "")).strip()
            if value:
                results.append(f"{term}: {value} [{source}]")
    return results
