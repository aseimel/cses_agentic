import tempfile
import unittest
from pathlib import Path

from src.matching.decision_engine import MatchingDecisionEngine, matching_category_summary
from src.matching.party_order import (
    ElectionResultsWorkbookParser,
    PartyOrderingRulesEngine,
    election_results_intake_summary,
    infer_election_context_from_macro_material,
    party_order_message,
    write_standardized_election_results_template,
)


class PartyOrderWorkflowTests(unittest.TestCase):
    def _write_workbook(self, path: Path) -> None:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Election Results"
        ws.cell(row=1, column=1, value="TESTLAND: 2024")
        ws.cell(row=3, column=1, value="Lower House")
        headers = ["Party label", "Numeric", "Party Name", "Votes", "% of Vote", "Seats", "% of Seats"]
        for index, header in enumerate(headers, start=1):
            ws.cell(row=5, column=index, value=header)
        rows = [
            ["PARTY A", 999001, "First Party", 1000, 40.0, 40, 40.0],
            ["PARTY B", 999002, "Second Party", 800, 32.0, 32, 32.0],
            ["PARTY C", 999003, "Third Party", 500, 20.0, 20, 20.0],
            ["", "", "Others", 200, 8.0, 8, 8.0],
        ]
        for row_index, row in enumerate(rows, start=6):
            for col_index, value in enumerate(row, start=1):
                ws.cell(row=row_index, column=col_index, value=value)
        wb.save(path)

    def test_parser_reads_standardized_election_results_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            study_dir = Path(folder)
            election_dir = study_dir / "Election Results"
            election_dir.mkdir()
            path = election_dir / "election_results.xlsx"
            self._write_workbook(path)

            tables = ElectionResultsWorkbookParser().parse_directory(study_dir)

            self.assertEqual(len(tables), 1)
            self.assertEqual(tables[0].election_context, "lower_house")
            self.assertEqual([party.code_letter for party in tables[0].parties[:3]], ["A", "B", "C"])
            self.assertEqual(tables[0].parties[0].numeric_code, "999001")

    def test_party_order_proposal_requires_joint_agreement_before_matching(self):
        with tempfile.TemporaryDirectory() as folder:
            study_dir = Path(folder)
            election_dir = study_dir / "Election Results"
            election_dir.mkdir()
            self._write_workbook(election_dir / "election_results.xlsx")

            tables = ElectionResultsWorkbookParser().parse_directory(study_dir)
            engine = PartyOrderingRulesEngine()
            proposal = engine.propose(tables, source_variables=["Q16a", "Q17a", "Q10LHb"])
            review_path, decision_path = engine.write_review(study_dir, proposal)

            self.assertTrue(review_path.exists())
            self.assertTrue(decision_path.exists())
            self.assertFalse(engine.is_approved(study_dir))
            self.assertEqual(proposal.status, "agreement_needed")
            self.assertEqual([party.party_name for party in proposal.proposed_parties[:3]], ["First Party", "Second Party", "Third Party"])
            self.assertIn("Micro processor and macro coder", party_order_message(proposal))

    def test_matching_engine_gates_party_variables_without_approved_order(self):
        decisions = MatchingDecisionEngine().decide(
            source_contexts=[{"name": "Q16a", "description": "Like-dislike Party A"}],
            matching_evidence={
                "target_variable_profiles": [
                    {"name": "F3018_A", "canonical_item_type": "party_vote_item", "section": "cses_module"}
                ]
            },
            party_order_approved=False,
            party_order_summary={"party_count": 3},
        )
        by_target = {item.target_variable: item for item in decisions}

        self.assertEqual(by_target["F3018_A"].status, "awaiting_party_order_agreement")
        summary = matching_category_summary(decisions, {"target_variable_profiles": [{"name": "F3018_A", "canonical_item_type": "party_vote_item"}]})
        self.assertGreater(summary["party_election_items"]["awaiting_party_ordering"], 0)

    def test_sweden_reference_workbook_party_order_is_reproduced_when_available(self):
        root = Path.cwd()
        workbook = root / "Sweden_2022" / "Election Results" / "SWE_2022_Election results.xlsx"
        if not workbook.exists():
            self.skipTest("Sweden election-results workbook is not available in this checkout")

        tables = ElectionResultsWorkbookParser().parse_file(workbook)
        proposal = PartyOrderingRulesEngine().propose(
            tables,
            context_hint=infer_election_context_from_macro_material(root / "Sweden_2022"),
        )
        observed = [(party.code_letter, party.numeric_code, party.party_name) for party in proposal.proposed_parties[:8]]

        self.assertEqual([item[0] for item in observed], list("ABCDEFGH"))
        self.assertEqual([item[1] for item in observed], ["752001", "752002", "752003", "752004", "752005", "752006", "752007", "752008"])
        self.assertIn("Social Democratic Party", observed[0][2])
        self.assertEqual(proposal.selected_context, "lower_house")

    def test_sweden_party_matching_after_party_order_approval_matches_reference_sources(self):
        from src.ingest.data_loader import DataLoader
        from src.matching.decision_engine import MatchingDecisionEngine, matching_category_summary

        root = Path.cwd()
        source = root / "Sweden_2022" / "micro" / "deposited datasets" / "CSES6_SWEDEN.dta"
        workbook = root / "Sweden_2022" / "Election Results" / "SWE_2022_Election results.xlsx"
        if not source.exists() or not workbook.exists():
            self.skipTest("Sweden reference files are not available in this checkout")

        dataset = DataLoader().load(source)
        source_contexts = [
            {"name": item.name, "description": item.description or "", "value_labels": item.value_labels or {}, "sample_values": item.sample_values or []}
            for item in dataset.variables.values()
        ]
        proposal = PartyOrderingRulesEngine().propose(
            ElectionResultsWorkbookParser().parse_file(workbook),
            source_variables=[item["name"] for item in source_contexts],
            context_hint=infer_election_context_from_macro_material(root / "Sweden_2022"),
        )
        decisions = MatchingDecisionEngine().decide(
            source_contexts=source_contexts,
            party_order_approved=True,
            party_order_summary={
                "party_count": len(proposal.proposed_parties),
                "selected_context": proposal.selected_context,
            },
        )
        by_target = {item.target_variable: item for item in decisions}
        expected_sources = {
            "F3011_LH_PL": "Q10LHb",
            "F3011_LH_PF": "Q10LHd",
            "F3016_LH_PL": "Q14b",
            "F3018_A": "Q16a",
            "F3018_H": "Q16h",
            "F3019_G": "Q17g1",
            "F3019_H": "Q17h",
            "F3019_I": "Q17g2",
            "F3020_A": "Q18a",
            "F3020_H": "Q18h",
            "F3023_1": "Q23a",
            "F3023_3": "Q23c",
        }
        for target, source_name in expected_sources.items():
            with self.subTest(target=target):
                self.assertEqual(by_target[target].source_variable, source_name)
                self.assertEqual(by_target[target].status, "proposed_match")

        generated = {
            "F3011_PR_1": "NOT_APPLICABLE_PRESIDENTIAL_ELECTION",
            "F3011_UH_PL": "NOT_APPLICABLE_UPPER_HOUSE_ELECTION",
            "F3011_LH_DC": "NOT_APPLICABLE_DISTRICT_CANDIDATE_VOTE",
            "F3018_I": "NO_APPROVED_PARTY_FOR_THIS_SLOT",
            "F3021_A": "OPTIONAL_ALTERNATIVE_SCALE_NOT_COLLECTED",
            "F3011_OUTGOV": "DERIVED_FROM_APPROVED_VOTE_CHOICE_AND_PARTY_METADATA",
            "F5000_A": "APPROVED_PARTY_ORDER",
            "F6000_LH_PL": "APPROVED_PARTY_ORDER",
        }
        for target, source_name in generated.items():
            with self.subTest(target=target):
                self.assertEqual(by_target[target].source_variable, source_name)
                self.assertIn(by_target[target].status, {"generated_from_party_context", "generated_from_party_order"})

        summary = matching_category_summary(decisions)
        self.assertEqual(summary["party_election_items"]["awaiting_party_ordering"], 0)
        self.assertEqual(summary["party_election_items"]["matched"], summary["party_election_items"]["total"])

    def test_intake_summary_detects_standardized_tables(self):
        with tempfile.TemporaryDirectory() as folder:
            study_dir = Path(folder)
            election_dir = study_dir / "Election Results"
            election_dir.mkdir()
            self._write_workbook(election_dir / "election_results.xlsx")

            summary = election_results_intake_summary(study_dir)

            self.assertEqual(summary["table_count"], 1)
            self.assertTrue(summary["standardized"])

    def test_election_results_discovery_uses_dedicated_folder_only(self):
        with tempfile.TemporaryDirectory() as folder:
            study_dir = Path(folder)
            (study_dir / "emails").mkdir()
            self._write_workbook(study_dir / "emails" / "election_results.xlsx")

            summary = election_results_intake_summary(study_dir)

            self.assertEqual(summary["table_count"], 0)
            self.assertFalse(summary["standardized"])

    def test_standardized_election_results_template_is_created_in_dedicated_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            study_dir = Path(folder)

            template = write_standardized_election_results_template(study_dir, ["raw_results.pdf"])

            self.assertTrue(template.exists())
            self.assertEqual(template.parent, study_dir / "Election Results")
            self.assertEqual(ElectionResultsWorkbookParser().parse_file(template), [])


if __name__ == "__main__":
    unittest.main()
